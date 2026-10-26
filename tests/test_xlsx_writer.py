from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from bi_etl.components.xlsx_reader import XLSXReader
from bi_etl.components.xlsx_writer import XLSXWriter

COLUMNS = ['id', 'name']


def make_writer(path, **kwargs) -> XLSXWriter:
    kwargs.setdefault('column_names', COLUMNS)
    return XLSXWriter(None, path, **kwargs)


def insert(writer: XLSXWriter, **values):
    return writer.insert(writer.Row(data=values))


def read_sheet(path, sheet_name=None) -> list:
    workbook = load_workbook(path)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def test_writes_header_and_rows(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
        insert(writer, id=2, name='b')
    assert read_sheet(path, 'Data') == [['id', 'name'], [1, 'a'], [2, 'b']]


def test_accepts_str_file_name(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(str(path)) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
    assert read_sheet(path) == [['id', 'name'], [1, 'a']]


def test_logical_name_defaults_to_file_basename(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        assert writer.logical_name == 'out.xlsx'


def test_explicit_logical_name(tmp_path):
    with make_writer(tmp_path / 'out.xlsx', logical_name='my_output') as writer:
        assert writer.logical_name == 'my_output'


def test_kwargs_set_attributes(tmp_path):
    with make_writer(tmp_path / 'out.xlsx', start_col=3, header_row=2) as writer:
        assert writer.start_col == 3
        assert writer.header_row == 2


def test_defaults(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        assert writer.write_only is True
        assert writer.read_only is False
        assert writer.headers_written is False
        assert writer.rows_inserted == 0
        assert writer.rows_inserted_this_sheet == 0


def test_close_without_touching_workbook_creates_no_file(tmp_path):
    path = tmp_path / 'out.xlsx'
    writer = make_writer(path)
    writer.close()
    assert writer.is_closed
    assert not path.exists()


def test_close_with_error_does_not_save(tmp_path):
    path = tmp_path / 'out.xlsx'
    writer = make_writer(path)
    writer.set_active_worksheet_by_name('Data')
    insert(writer, id=1, name='a')
    writer.close(error=True)
    assert writer.is_closed
    assert not path.exists()


def test_obtain_column_names_raises_when_not_set(tmp_path):
    with XLSXWriter(None, tmp_path / 'out.xlsx') as writer:
        with pytest.raises(ValueError, match='Column names must be explicitly set'):
            _ = writer.column_names


def test_write_header_requires_column_names(tmp_path):
    with XLSXWriter(None, tmp_path / 'out.xlsx') as writer:
        writer.set_active_worksheet_by_name('Data')
        with pytest.raises(ValueError):
            writer.write_header()


def test_write_header_manual_then_insert_does_not_duplicate_header(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        writer.write_header()
        assert writer.headers_written
        insert(writer, id=1, name='a')
    assert read_sheet(path) == [['id', 'name'], [1, 'a']]


@pytest.mark.parametrize('write_only', [True, False])
def test_header_row_positions_header(tmp_path, write_only):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, write_only=write_only, header_row=3, start_col=2) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
    assert read_sheet(path) == [
        [None, None, None],
        [None, None, None],
        [None, 'id', 'name'],
        [None, 1, 'a'],
    ]


@pytest.mark.parametrize('write_only', [True, False])
def test_header_row_default_adds_no_blank_rows(tmp_path, write_only):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, write_only=write_only) as writer:
        writer.set_active_worksheet_by_name('Data')
        writer.write_header()
    assert read_sheet(path) == [['id', 'name']]


def test_make_table_respects_header_row(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, write_only=False, header_row=3) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
        writer.make_table_from_inserted_data()
    workbook = load_workbook(path)
    try:
        assert workbook['Data'].tables['Tbl_Data'].ref == 'A3:B4'
    finally:
        workbook.close()


def test_insert_counts_rows(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
        insert(writer, id=2, name='b')
        assert writer.rows_inserted == 2
        assert writer.rows_inserted_this_sheet == 2


def test_insert_single_row_returns_new_row(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        writer.set_active_worksheet_by_name('Data')
        new_row = writer.insert_row(writer.Row(data={'id': 1, 'name': 'a'}))
        assert new_row['id'] == 1
        assert new_row['name'] == 'a'


def test_insert_list_writes_all_rows(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        writer.insert([
            writer.Row(data={'id': 1, 'name': 'a'}),
            writer.Row(data={'id': 2, 'name': 'b'}),
        ])
        assert writer.rows_inserted == 2
    assert read_sheet(path) == [['id', 'name'], [1, 'a'], [2, 'b']]


def test_additional_insert_values_add_and_override(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        writer.insert(
            writer.Row(data={'id': 1, 'name': 'a'}),
            additional_insert_values={'name': 'override'},
        )
        new_row = writer.insert_row(
            writer.Row(data={'id': 2, 'name': 'b'}),
            additional_insert_values={'name': 'override'},
        )
        assert new_row['name'] == 'override'
    assert read_sheet(path) == [['id', 'name'], [1, 'override'], [2, 'override']]


def test_extra_source_columns_are_dropped(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a', extra='ignored')
    assert read_sheet(path) == [['id', 'name'], [1, 'a']]


def test_missing_columns_written_blank(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1)
    assert read_sheet(path) == [['id', 'name'], [1, None]]


def test_empty_source_row_asserts(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        writer.set_active_worksheet_by_name('Data')
        with pytest.raises(AssertionError, match='Empty row passed'):
            writer.insert_row(writer.Row())


def test_start_col_offsets_header_and_values(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, start_col=3) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
    assert read_sheet(path) == [[None, None, 'id', 'name'], [None, None, 1, 'a']]


def test_start_col_offsets_data_values(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, start_col=3) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
    assert read_sheet(path)[1] == [None, None, 1, 'a']


def test_value_types_round_trip(tmp_path):
    path = tmp_path / 'out.xlsx'
    moment = datetime(2020, 1, 2, 3, 4, 5)
    with make_writer(path, column_names=['i', 'f', 's', 'd', 'b']) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, i=1, f=2.5, s='text', d=moment, b=True)
    assert read_sheet(path)[1] == [1, 2.5, 'text', moment, True]


def test_illegal_xml_characters_are_replaced(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='bad\x00\x08value\x1f')
    # Pins current behavior: the replacement text is a literal backslash followed by '?'
    assert read_sheet(path)[1] == [1, 'bad\\?\\?value\\?']


def test_legal_control_characters_are_kept(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='line1\nline2\ttab')
    assert read_sheet(path)[1] == [1, 'line1\nline2\ttab']


def test_set_active_worksheet_by_name_creates_sheet(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        writer.set_active_worksheet_by_name('New')
        assert writer.active_worksheet_name == 'New'
        assert writer.get_sheet_names() == ['New']


def test_multiple_sheets_each_get_header_and_counts(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('One')
        insert(writer, id=1, name='a')
        insert(writer, id=2, name='b')
        assert writer.rows_inserted_this_sheet == 2

        writer.set_active_worksheet_by_name('Two')
        assert writer.headers_written is False
        assert writer.rows_inserted_this_sheet == 0
        writer.column_names = ['x']
        insert(writer, x=9)
        assert writer.rows_inserted_this_sheet == 1
        assert writer.rows_inserted == 3
    assert read_sheet(path, 'One') == [['id', 'name'], [1, 'a'], [2, 'b']]
    assert read_sheet(path, 'Two') == [['x'], [9]]


def test_set_active_worksheet_by_number_missing_sheet_message(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        writer.set_active_worksheet_by_name('One')
        with pytest.raises(ValueError, match='Use set_active_worksheet_by_name to create new sheets'):
            writer.set_active_worksheet_by_number(5)


def test_set_active_worksheet_by_number_existing_sheet(tmp_path):
    with make_writer(tmp_path / 'out.xlsx', write_only=False) as writer:
        writer.set_active_worksheet_by_name('One')
        writer.set_active_worksheet_by_name('Two')
        writer.set_active_worksheet_by_number(0)
        assert writer.active_worksheet_name == 'One'


def test_set_widths(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, write_only=False) as writer:
        writer.set_active_worksheet_by_name('Data')
        writer.set_widths([10, 20.5])
        insert(writer, id=1, name='a')
    workbook = load_workbook(path)
    try:
        dims = workbook['Data'].column_dimensions
        assert dims['A'].width == 10
        assert dims['B'].width == 20.5
    finally:
        workbook.close()


def test_set_columns_and_widths(tmp_path):
    path = tmp_path / 'out.xlsx'
    with XLSXWriter(None, path, write_only=False) as writer:
        writer.set_active_worksheet_by_name('Data')
        writer.set_columns_and_widths({'id': 8, 'name': 30})
        assert writer.column_names == ['id', 'name']
        insert(writer, id=1, name='a')
    workbook = load_workbook(path)
    try:
        dims = workbook['Data'].column_dimensions
        assert dims['A'].width == 8
        assert dims['B'].width == 30
    finally:
        workbook.close()


def test_make_table_from_inserted_data(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, write_only=False) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
        insert(writer, id=2, name='b')
        writer.make_table_from_inserted_data(style_name='TableStyleLight1')
    workbook = load_workbook(path)
    try:
        tables = workbook['Data'].tables
        assert list(tables.keys()) == ['Tbl_Data']
        table = tables['Tbl_Data']
        assert table.ref == 'A1:B3'
        assert [c.name for c in table.tableColumns] == COLUMNS
        assert table.tableStyleInfo.name == 'TableStyleLight1'
        assert table.autoFilter.ref == 'A1:B3'
    finally:
        workbook.close()


def test_make_table_respects_start_col(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, write_only=False, start_col=3) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
        writer.make_table_from_inserted_data()
    workbook = load_workbook(path)
    try:
        assert workbook['Data'].tables['Tbl_Data'].ref == 'C1:D2'
    finally:
        workbook.close()


def test_make_table_without_data_is_noop(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path, write_only=False) as writer:
        writer.set_active_worksheet_by_name('Data')
        writer.make_table_from_inserted_data()
    workbook = load_workbook(path)
    try:
        assert len(workbook['Data'].tables) == 0
    finally:
        workbook.close()


def test_write_only_workbook_is_write_only(tmp_path):
    with make_writer(tmp_path / 'out.xlsx') as writer:
        assert writer.workbook.write_only is True


def test_not_write_only_creates_new_workbook_without_default_sheet(tmp_path):
    with make_writer(tmp_path / 'out.xlsx', write_only=False) as writer:
        assert writer.workbook.write_only is False
        assert writer.get_sheet_names() == []
        writer.set_active_worksheet_by_name('Data')
        assert writer.get_sheet_names() == ['Data']


def test_not_write_only_appends_to_existing_file(tmp_path):
    path = tmp_path / 'out.xlsx'
    existing = Workbook()
    existing.active.title = 'Data'
    existing.active.append(['id', 'name'])
    existing.active.append([1, 'a'])
    existing.save(path)
    existing.close()

    with make_writer(path, write_only=False) as writer:
        writer.set_active_worksheet_by_name('Data')
        # Switching to an existing sheet resets column names
        writer.column_names = COLUMNS
        # Header is already present in the file, so skip writing it again
        writer.headers_written = True
        insert(writer, id=2, name='b')
    assert read_sheet(path, 'Data') == [['id', 'name'], [1, 'a'], [2, 'b']]


def test_not_write_only_existing_sheet_resets_column_names(tmp_path):
    path = tmp_path / 'out.xlsx'
    existing = Workbook()
    existing.active.title = 'Data'
    existing.save(path)
    existing.close()

    with make_writer(path, write_only=False) as writer:
        writer.set_active_worksheet_by_name('Data')
        assert writer._column_names is None
        assert writer.rows_inserted_this_sheet == 0


def test_output_readable_by_xlsx_reader(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
        insert(writer, id=2, name='b')
    with XLSXReader(None, path) as reader:
        rows = [(row['id'], row['name']) for row in reader]
    assert rows == [(1, 'a'), (2, 'b')]


def test_context_manager_closes_and_saves(tmp_path):
    path = tmp_path / 'out.xlsx'
    with make_writer(path) as writer:
        writer.set_active_worksheet_by_name('Data')
        insert(writer, id=1, name='a')
        assert not path.exists()
    assert writer.is_closed
    assert path.exists()
